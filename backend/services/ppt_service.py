# backend/services/ppt_service.py

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import re

class PPTService:

    def __init__(self):
        pass

    def clean_markdown(self, markdown_text: str):
        """
        Converts markdown to structured slides.
        Returns a list of dicts: [{"title": str, "bullets": [str, ...]}, ...]
        """
        slides = []
        current_slide = None

        for line in markdown_text.split("\n"):
            line = line.strip()
            if not line:
                continue

            # Slide title -> # Heading1
            if line.startswith("# "):
                if current_slide:
                    slides.append(current_slide)
                current_slide = {"title": line[2:], "bullets": []}

            # Bullet points -> - or * 
            elif line.startswith("- ") or line.startswith("* "):
                if current_slide:
                    current_slide["bullets"].append(line[2:])

        # Add last slide
        if current_slide:
            slides.append(current_slide)

        return slides

    def markdown_to_ppt(self, markdown_text: str) -> bytes:
        prs = Presentation()
        prs.slide_width = Inches(13.33)   # widescreen 16:9
        prs.slide_height = Inches(7.5)

        slides_data = self.clean_markdown(markdown_text)

        for s in slides_data:
            slide_layout = prs.slide_layouts[1]  # Title + Content
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = s["title"]

            # Add bullets
            content = slide.placeholders[1]
            tf = content.text_frame
            tf.clear()
            for b in s["bullets"]:
                p = tf.add_paragraph()
                p.text = b
                p.font.size = Pt(24)
                p.level = 0
                p.alignment = PP_ALIGN.LEFT

        # Save to bytes
        from io import BytesIO
        ppt_buffer = BytesIO()
        prs.save(ppt_buffer)
        ppt_bytes = ppt_buffer.getvalue()
        ppt_buffer.close()
        return ppt_bytes